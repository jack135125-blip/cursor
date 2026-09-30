# -*- coding: utf-8 -*-
"""
창원여고 주간 업무 계획 - 오늘의 알림 팝업
Google 스프레드시트에서 오늘 날짜 열만 읽어 둥근 팝업으로 표시합니다.
창을 닫으면 프로그램이 종료됩니다.
"""

from __future__ import annotations

import csv
import ctypes
import io
import json
import re
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
from calendar import monthrange
from datetime import date, datetime
from pathlib import Path

# customtkinter import 전에 DPI를 선언해야 흐림이 줄어듭니다.
def _prepare_windows_dpi() -> float:
    if sys.platform != "win32":
        return 1.0
    try:
        ctypes.windll.shcore.SetProcessDpiAwareness(2)  # Per-monitor DPI aware
    except Exception:
        try:
            ctypes.windll.user32.SetProcessDPIAware()
        except Exception:
            pass
    try:
        dpi = float(ctypes.windll.user32.GetDpiForSystem())
        return max(dpi / 96.0, 1.0)
    except Exception:
        return 1.0


DPI_SCALE = _prepare_windows_dpi()

import customtkinter as ctk

try:
    # 자동 추가 스케일링과 OS DPI가 겹치면 글씨가 번져 보입니다.
    ctk.deactivate_automatic_dpi_awareness()
    ctk.set_widget_scaling(1.0)
    ctk.set_window_scaling(1.0)
except Exception:
    pass

# ---------------------------------------------------------------------------
# 설정
# ---------------------------------------------------------------------------
SPREADSHEET_ID = "1dOzQRmT3eIKW3OuIQlDAw7WkHRqn8j-kdLbv18x2zPM"
CACHE_PATH = Path(__file__).with_name("today_schedule_cache.json")

# 아기자기한 파스텔 톤 (민트 + 코랄)
COLORS = {
    "bg": "#E8F6F3",
    "card": "#FFFFFF",
    "header": "#5EB7A8",
    "header_text": "#FFFFFF",
    "chip": "#FFE8E1",
    "chip_text": "#C45C48",
    "dept": "#2F5D56",
    "body": "#3A4A47",
    "muted": "#7A8F8A",
    "button": "#5EB7A8",
    "button_hover": "#4AA496",
    "empty": "#F7FBFA",
    "shadow": "#D5EBE6",
}

WEEKDAY_KR = ("월요일", "화요일", "수요일", "목요일", "금요일", "토요일", "일요일")
USER_AGENT = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) TodaySchedulePopup/1.0"

# 배달의민족 주아 고정 (없으면 대체 글씨체)
_FONT_CANDIDATES = (
    "배달의민족 주아",
    "나눔스퀘어라운드 Regular",
    "나눔스퀘어라운드",
    "맑은 고딕",
)

FONT_UI = "배달의민족 주아"
FONT_UI_BOLD = "배달의민족 주아"


def _resolve_fonts() -> None:
    global FONT_UI, FONT_UI_BOLD
    try:
        from tkinter import font as tkfont

        available = set(tkfont.families())
        for name in _FONT_CANDIDATES:
            if name in available:
                FONT_UI = name
                FONT_UI_BOLD = name
                break
    except Exception:
        FONT_UI = "맑은 고딕"
        FONT_UI_BOLD = "맑은 고딕"


def ui_font(size: int, *, emphasize: bool = False):
    """주아는 이미 두꺼운 서체라 가짜 bold를 쓰면 획이 뭉칩니다."""
    # emphasize여도 weight는 normal 유지 (주아/라운드체 가독성)
    _ = emphasize
    scaled = max(1, int(round(size * 1.10)))
    return ctk.CTkFont(family=FONT_UI, size=scaled, weight="normal")


# 드롭다운 현재 선택 행 강조색
DROPDOWN_SELECTED_BG = "#FFE0D6"
DROPDOWN_SELECTED_FG = "#C45C48"
DROPDOWN_SELECTED_HOVER = "#FFCFC2"


def attach_dropdown_selection_highlight(option_menu: ctk.CTkOptionMenu, get_current) -> None:
    """열린 목록에서 현재 선택된 값을 다른 색/체크로 구분합니다."""
    dropdown = option_menu._dropdown_menu

    def _rebuild_commands() -> None:
        dropdown.delete(0, "end")
        values = dropdown._values or []
        current = get_current()
        min_width = getattr(dropdown, "_min_character_width", 18)

        for value in values:
            selected = value == current
            mark = "✓ " if selected else "   "
            label = f"{mark}{value}".ljust(min_width + 3)
            kwargs = {
                "label": label,
                "command": (lambda v=value: dropdown._button_callback(v)),
                "compound": "left",
            }
            # Windows tk Menu는 항목별 배경색을 지원합니다.
            if sys.platform.startswith("win"):
                if selected:
                    kwargs.update(
                        {
                            "background": DROPDOWN_SELECTED_BG,
                            "foreground": DROPDOWN_SELECTED_FG,
                            "activebackground": DROPDOWN_SELECTED_HOVER,
                            "activeforeground": DROPDOWN_SELECTED_FG,
                        }
                    )
            dropdown.add_command(**kwargs)

    dropdown._add_menu_commands = _rebuild_commands
    original_open = dropdown.open

    def _open_with_highlight(x, y):
        _rebuild_commands()
        original_open(x, y)

    dropdown.open = _open_with_highlight
    _rebuild_commands()


# ---------------------------------------------------------------------------
# 스프레드시트 연동
# ---------------------------------------------------------------------------
# gviz CSV는 일부 칸(특히 여러 날에 걸친 일정)을 비워서 내보내므로
# xlsx보내기로 읽어 옵니다.
_workbook_rows_cache: dict[str, list[list[str]]] | None = None


def _candidate_sheet_names() -> list[str]:
    names = [f"2학기 {i}주" for i in range(1, 23)]
    names += [f"{i}주" for i in range(1, 22)]
    return names


def _load_cache() -> dict:
    if not CACHE_PATH.exists():
        return {}
    try:
        return json.loads(CACHE_PATH.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return {}


def _save_cache(sheet_name: str) -> None:
    try:
        CACHE_PATH.write_text(
            json.dumps({"last_sheet": sheet_name}, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
    except OSError:
        pass


def _download_workbook_bytes() -> bytes:
    url = f"https://docs.google.com/spreadsheets/d/{SPREADSHEET_ID}/export?format=xlsx"
    req = urllib.request.Request(url, headers={"User-Agent": USER_AGENT})
    with urllib.request.urlopen(req, timeout=60) as resp:
        return resp.read()


def _cell_to_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return (
            f"{value.year}년 {value.month}월 {value.day}일 "
            f"{WEEKDAY_KR[value.weekday()]}"
        )
    if isinstance(value, date) and not isinstance(value, datetime):
        return (
            f"{value.year}년 {value.month}월 {value.day}일 "
            f"{WEEKDAY_KR[value.weekday()]}"
        )
    return str(value)


def _worksheet_to_rows(ws) -> list[list[str]]:
    """병합 셀 값을 범위 전체에 채운 뒤 행 목록으로 변환합니다."""
    merge_values: dict[tuple[int, int], str] = {}
    for merged in ws.merged_cells.ranges:
        text = _cell_to_text(ws.cell(merged.min_row, merged.min_col).value)
        for r in range(merged.min_row, merged.max_row + 1):
            for c in range(merged.min_col, merged.max_col + 1):
                merge_values[(r, c)] = text

    max_row = min(ws.max_row or 0, 40)
    max_col = min(ws.max_column or 0, 12)
    rows: list[list[str]] = []
    for r in range(1, max_row + 1):
        row: list[str] = []
        for c in range(1, max_col + 1):
            if (r, c) in merge_values:
                row.append(merge_values[(r, c)])
            else:
                row.append(_cell_to_text(ws.cell(r, c).value))
        rows.append(row)
    return rows


def load_all_sheet_rows(force_refresh: bool = False) -> dict[str, list[list[str]]]:
    global _workbook_rows_cache
    if _workbook_rows_cache is not None and not force_refresh:
        return _workbook_rows_cache

    import openpyxl

    data = _download_workbook_bytes()
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    sheets: dict[str, list[list[str]]] = {}
    for name in wb.sheetnames:
        sheets[name] = _worksheet_to_rows(wb[name])
    wb.close()
    _workbook_rows_cache = sheets
    return sheets


def fetch_sheet_csv(sheet_name: str | None = None, gid: str | None = None) -> list[list[str]]:
    """하위 호환용. 가능하면 xlsx 캐시에서 읽고, 없으면 export CSV(gid)로 읽습니다."""
    if sheet_name:
        sheets = load_all_sheet_rows()
        if sheet_name in sheets:
            return sheets[sheet_name]

    if gid:
        url = (
            f"https://docs.google.com/spreadsheets/d/{SPREADSHEET_ID}/"
            f"export?format=csv&gid={gid}"
        )
        req = urllib.request.Request(url, headers={"User-Agent": USER_AGENT})
        with urllib.request.urlopen(req, timeout=30) as resp:
            raw = resp.read().decode("utf-8-sig", errors="replace")
        return list(csv.reader(io.StringIO(raw)))

    raise FileNotFoundError(f"시트를 찾지 못했습니다: sheet={sheet_name!r} gid={gid!r}")


def clean_cell(value: str) -> str:
    text = (value or "").replace("\r", "\n").strip()
    text = re.sub(r"\n+", " ", text)
    text = re.sub(r"[ \t]+", " ", text)
    return text.strip(" -")


def split_tasks(text: str) -> list[str]:
    text = clean_cell(text)
    if not text:
        return []
    # "-항목 -항목" 형태 우선 분리
    parts = [p.strip() for p in re.split(r"\s*-\s*", text) if p.strip()]
    if len(parts) > 1:
        return parts
    return [text]


def find_date_column(header_row: list[str], target: date) -> int | None:
    needle = f"{target.year}년 {target.month}월 {target.day}일"
    for idx, cell in enumerate(header_row):
        if needle in clean_cell(cell):
            return idx
    return None


def find_header_row(rows: list[list[str]], target: date) -> tuple[int, int] | None:
    needle = f"{target.year}년 {target.month}월 {target.day}일"
    for r_idx, row in enumerate(rows[:8]):
        for c_idx, cell in enumerate(row):
            if needle in clean_cell(cell):
                return r_idx, c_idx
    return None


def extract_today_items(rows: list[list[str]], target: date) -> tuple[str, list[tuple[str, list[str]]]] | None:
    found = find_header_row(rows, target)
    if not found:
        return None
    header_idx, col_idx = found
    header_label = clean_cell(rows[header_idx][col_idx])

    items: list[tuple[str, list[str]]] = []
    for row in rows[header_idx + 1 :]:
        if not row:
            continue
        dept = clean_cell(row[0] if len(row) > 0 else "")
        if not dept or "부서" in dept:
            continue
        if dept in {"토의 안건"}:
            # 안건도 보여 주되, 내용 없으면 생략
            pass
        cell = row[col_idx] if col_idx < len(row) else ""
        tasks = split_tasks(cell)
        if tasks:
            items.append((dept, tasks))
    return header_label, items


def ordered_sheet_candidates(available: list[str] | None = None) -> list[str]:
    cache = _load_cache()
    last = cache.get("last_sheet")
    base = available if available is not None else _candidate_sheet_names()
    ordered: list[str] = []
    if last and last in base:
        ordered.append(last)
        i = base.index(last)
        for j in (i - 1, i + 1, i - 2, i + 2):
            if 0 <= j < len(base) and base[j] not in ordered:
                ordered.append(base[j])
    for name in base:
        if name not in ordered:
            ordered.append(name)
    return ordered


def load_today_schedule(target: date | None = None) -> dict:
    target = target or date.today()
    errors: list[str] = []

    try:
        sheets = load_all_sheet_rows()
    except (urllib.error.URLError, TimeoutError, OSError, ValueError) as exc:
        return {
            "ok": False,
            "sheet_name": None,
            "date_label": None,
            "items": [],
            "target": target,
            "error": f"스프레드시트를 불러오지 못했습니다.\n{exc}",
        }

    # 실제 탭 이름을 우선 사용하고, 예전 후보 이름도 함께 시도
    available = list(sheets.keys())
    for sheet_name in ordered_sheet_candidates(available):
        rows = sheets.get(sheet_name)
        if rows is None:
            continue
        try:
            parsed = extract_today_items(rows, target)
        except Exception as exc:  # noqa: BLE001
            errors.append(f"{sheet_name}: {exc}")
            continue
        if not parsed:
            continue

        header_label, items = parsed
        _save_cache(sheet_name)
        return {
            "ok": True,
            "sheet_name": sheet_name,
            "date_label": header_label,
            "items": items,
            "target": target,
        }

    return {
        "ok": False,
        "sheet_name": None,
        "date_label": None,
        "items": [],
        "target": target,
        "error": "선택한 날짜가 들어 있는 주간 시트를 찾지 못했습니다."
        + (("\n" + "\n".join(errors[:3])) if errors else ""),
    }


def year_options(center: date | None = None) -> list[str]:
    center = center or date.today()
    years = sorted({center.year - 1, center.year, center.year + 1, 2026})
    return [f"{y}년" for y in years]


def month_options() -> list[str]:
    return [f"{m}월" for m in range(1, 13)]


def day_options(year: int, month: int) -> list[str]:
    last = monthrange(year, month)[1]
    return [f"{d}일" for d in range(1, last + 1)]


def parse_part(value: str) -> int:
    return int(re.sub(r"[^0-9]", "", value))


# ---------------------------------------------------------------------------
# UI
# ---------------------------------------------------------------------------
class TodaySchedulePopup(ctk.CTk):
    def __init__(self, payload: dict):
        super().__init__()
        self.payload = payload
        self._drag_x = 0
        self._drag_y = 0
        self._loading = False
        self._updating_menus = False
        self.body_frame: ctk.CTkScrollableFrame | None = None
        self.sheet_hint: ctk.CTkLabel | None = None
        self.weekday_label: ctk.CTkLabel | None = None
        self.year_menu: ctk.CTkOptionMenu | None = None
        self.month_menu: ctk.CTkOptionMenu | None = None
        self.day_menu: ctk.CTkOptionMenu | None = None

        ctk.set_appearance_mode("light")
        ctk.set_default_color_theme("green")
        _resolve_fonts()

        self.title("오늘의 주간 업무")
        w = int(680 * DPI_SCALE)
        h = int(860 * DPI_SCALE)
        self.geometry(f"{w}x{h}")
        self.minsize(int(620 * DPI_SCALE), int(720 * DPI_SCALE))
        self.configure(fg_color=COLORS["bg"])
        self.attributes("-topmost", True)

        # 둥근 카드 느낌의 외곽
        self.outer = ctk.CTkFrame(
            self,
            fg_color=COLORS["shadow"],
            corner_radius=30,
            border_width=0,
        )
        self.outer.pack(fill="both", expand=True, padx=16, pady=16)

        self.card = ctk.CTkFrame(
            self.outer,
            fg_color=COLORS["card"],
            corner_radius=26,
            border_width=0,
        )
        self.card.pack(fill="both", expand=True, padx=7, pady=7)

        self._build_header()
        self._build_body_container()
        self._render_body()
        self._build_footer()

        self.bind("<Escape>", lambda _e: self.close_app())
        self.protocol("WM_DELETE_WINDOW", self.close_app)

        # 화면 중앙
        self.after(30, self._center_window)

    def _center_window(self) -> None:
        self.update_idletasks()
        w, h = self.winfo_width(), self.winfo_height()
        sw, sh = self.winfo_screenwidth(), self.winfo_screenheight()
        x, y = max((sw - w) // 2, 0), max((sh - h) // 5, 40)
        self.geometry(f"{w}x{h}+{x}+{y}")

    def _menu_style(self) -> dict:
        return {
            "height": 40,
            "corner_radius": 14,
            "fg_color": "#FFFFFF",
            "button_color": "#FFE8E1",
            "button_hover_color": "#FFD5CA",
            "dropdown_fg_color": "#FFFFFF",
            "dropdown_hover_color": "#E8F6F3",
            "dropdown_text_color": COLORS["dept"],
            "text_color": COLORS["dept"],
            "font": ui_font(16),
            "dropdown_font": ui_font(15),
            "anchor": "center",
        }

    def _build_header(self) -> None:
        header = ctk.CTkFrame(self.card, fg_color=COLORS["header"], corner_radius=22)
        header.pack(fill="x", padx=20, pady=(20, 14))
        header.bind("<ButtonPress-1>", self._start_drag)
        header.bind("<B1-Motion>", self._on_drag)

        target: date = self.payload.get("target") or date.today()

        top = ctk.CTkFrame(header, fg_color="transparent")
        top.pack(fill="x", padx=20, pady=(16, 8))

        ctk.CTkLabel(
            top,
            text="오늘의 주간 업무 알림",
            font=ui_font(19, emphasize=True),
            text_color=COLORS["header_text"],
            anchor="w",
        ).pack(side="left")

        picker = ctk.CTkFrame(header, fg_color="transparent")
        picker.pack(fill="x", padx=20, pady=(0, 8))

        style = self._menu_style()
        self.year_var = ctk.StringVar(value=f"{target.year}년")
        self.month_var = ctk.StringVar(value=f"{target.month}월")
        self.day_var = ctk.StringVar(value=f"{target.day}일")

        self.year_menu = ctk.CTkOptionMenu(
            picker,
            values=year_options(target),
            variable=self.year_var,
            command=self._on_year_or_month_changed,
            width=110,
            **style,
        )
        self.year_menu.pack(side="left", padx=(0, 8))
        attach_dropdown_selection_highlight(self.year_menu, self.year_var.get)

        self.month_menu = ctk.CTkOptionMenu(
            picker,
            values=month_options(),
            variable=self.month_var,
            command=self._on_year_or_month_changed,
            width=96,
            **style,
        )
        self.month_menu.pack(side="left", padx=(0, 8))
        attach_dropdown_selection_highlight(self.month_menu, self.month_var.get)

        self.day_menu = ctk.CTkOptionMenu(
            picker,
            values=day_options(target.year, target.month),
            variable=self.day_var,
            command=self._on_day_changed,
            width=96,
            **style,
        )
        self.day_menu.pack(side="left", padx=(0, 10))
        attach_dropdown_selection_highlight(self.day_menu, self.day_var.get)

        self.weekday_label = ctk.CTkLabel(
            picker,
            text=WEEKDAY_KR[target.weekday()],
            font=ui_font(16),
            text_color="#F3FFFC",
            anchor="w",
        )
        self.weekday_label.pack(side="left")

        ctk.CTkFrame(header, fg_color="transparent", height=8).pack()

    def _selected_date(self) -> date | None:
        try:
            y = parse_part(self.year_var.get())
            m = parse_part(self.month_var.get())
            d = parse_part(self.day_var.get())
            last = monthrange(y, m)[1]
            d = min(d, last)
            return date(y, m, d)
        except (ValueError, AttributeError):
            return None

    def _sync_day_menu(self) -> None:
        if not self.day_menu:
            return
        try:
            y = parse_part(self.year_var.get())
            m = parse_part(self.month_var.get())
        except ValueError:
            return
        days = day_options(y, m)
        current_day = self.day_var.get()
        if current_day not in days:
            current_day = days[-1]
            self.day_var.set(current_day)
        self._updating_menus = True
        self.day_menu.configure(values=days)
        self.day_menu.set(current_day)
        self._updating_menus = False

    def _on_year_or_month_changed(self, _value: str) -> None:
        if self._updating_menus or self._loading:
            return
        self._sync_day_menu()
        self._reload_selected_date()

    def _on_day_changed(self, _value: str) -> None:
        if self._updating_menus or self._loading:
            return
        self._reload_selected_date()

    def _reload_selected_date(self) -> None:
        target = self._selected_date()
        if not target:
            return
        current = self.payload.get("target")
        if current == target:
            if self.weekday_label:
                self.weekday_label.configure(text=WEEKDAY_KR[target.weekday()])
            return

        if self.weekday_label:
            self.weekday_label.configure(text=WEEKDAY_KR[target.weekday()])

        self._loading = True
        self._render_body()
        self.update_idletasks()

        self.payload = load_today_schedule(target)
        self._loading = False
        self._render_body()
        self._update_sheet_hint()

    def _build_body_container(self) -> None:
        self.body_frame = ctk.CTkScrollableFrame(
            self.card,
            fg_color=COLORS["empty"],
            corner_radius=20,
            scrollbar_button_color=COLORS["header"],
            scrollbar_button_hover_color=COLORS["button_hover"],
        )
        self.body_frame.pack(fill="both", expand=True, padx=18, pady=(0, 10))

    def _clear_body(self) -> None:
        if not self.body_frame:
            return
        for child in self.body_frame.winfo_children():
            child.destroy()

    def _render_body(self) -> None:
        self._clear_body()
        body = self.body_frame
        if body is None:
            return

        if self._loading:
            self._add_message(body, "일정을 불러오는 중…", "잠시만 기다려 주세요.")
            return

        if not self.payload.get("ok"):
            self._add_message(
                body,
                "일정을 불러오지 못했어요",
                self.payload.get("error") or "잠시 후 다시 시도해 주세요.",
            )
            return

        items: list[tuple[str, list[str]]] = self.payload.get("items") or []
        if not items:
            self._add_message(
                body,
                "선택한 날짜에 등록된 업무가 없어요",
                "주간 업무 계획표의 해당 날짜 열이 비어 있습니다.",
            )
            return

        count_chip = ctk.CTkLabel(
            body,
            text=f"  부서 {len(items)}곳의 일정  ",
            font=ui_font(15, emphasize=True),
            text_color=COLORS["chip_text"],
            fg_color=COLORS["chip"],
            corner_radius=14,
        )
        count_chip.pack(anchor="w", padx=10, pady=(12, 8))

        for dept, tasks in items:
            self._add_dept_card(body, dept, tasks)

    def _add_message(self, parent, title: str, detail: str) -> None:
        box = ctk.CTkFrame(parent, fg_color=COLORS["card"], corner_radius=20)
        box.pack(fill="x", padx=10, pady=18)
        ctk.CTkLabel(
            box,
            text=title,
            font=ui_font(19, emphasize=True),
            text_color=COLORS["dept"],
        ).pack(padx=18, pady=(20, 6))
        ctk.CTkLabel(
            box,
            text=detail,
            font=ui_font(16),
            text_color=COLORS["muted"],
            wraplength=540,
            justify="left",
        ).pack(padx=18, pady=(0, 20))

    def _add_dept_card(self, parent, dept: str, tasks: list[str]) -> None:
        card = ctk.CTkFrame(parent, fg_color=COLORS["card"], corner_radius=20)
        card.pack(fill="x", padx=10, pady=7)

        badge = ctk.CTkLabel(
            card,
            text=f"  {dept}  ",
            font=ui_font(16, emphasize=True),
            text_color=COLORS["chip_text"],
            fg_color=COLORS["chip"],
            corner_radius=14,
            anchor="w",
        )
        badge.pack(anchor="w", padx=16, pady=(14, 8))

        for task in tasks:
            row = ctk.CTkFrame(card, fg_color="transparent")
            row.pack(fill="x", padx=16, pady=3)
            ctk.CTkLabel(
                row,
                text="●",
                font=ctk.CTkFont(size=12),
                text_color=COLORS["header"],
                width=20,
            ).pack(side="left", anchor="n", pady=4)
            ctk.CTkLabel(
                row,
                text=task,
                font=ui_font(16),
                text_color=COLORS["body"],
                wraplength=540,
                justify="left",
                anchor="w",
            ).pack(side="left", fill="x", expand=True)

        ctk.CTkFrame(card, fg_color="transparent", height=12).pack()

    def _build_footer(self) -> None:
        footer = ctk.CTkFrame(self.card, fg_color="transparent")
        footer.pack(fill="x", padx=18, pady=(4, 18))

        sheet = self.payload.get("sheet_name")
        hint = f"시트: {sheet}" if sheet else "Google 스프레드시트 연동"
        self.sheet_hint = ctk.CTkLabel(
            footer,
            text=hint,
            font=ui_font(14),
            text_color=COLORS["muted"],
            anchor="w",
        )
        self.sheet_hint.pack(side="left")

        ctk.CTkButton(
            footer,
            text="확인했어요",
            width=150,
            height=48,
            corner_radius=24,
            fg_color=COLORS["button"],
            hover_color=COLORS["button_hover"],
            text_color="white",
            font=ui_font(17, emphasize=True),
            command=self.close_app,
        ).pack(side="right")

    def _update_sheet_hint(self) -> None:
        if not self.sheet_hint:
            return
        sheet = self.payload.get("sheet_name")
        hint = f"시트: {sheet}" if sheet else "Google 스프레드시트 연동"
        self.sheet_hint.configure(text=hint)

    def _start_drag(self, event) -> None:
        self._drag_x = event.x_root - self.winfo_x()
        self._drag_y = event.y_root - self.winfo_y()

    def _on_drag(self, event) -> None:
        x = event.x_root - self._drag_x
        y = event.y_root - self._drag_y
        self.geometry(f"+{x}+{y}")

    def close_app(self) -> None:
        self.destroy()


def show_loading_then_popup() -> None:
    # 짧은 로딩 창
    boot = ctk.CTk()
    ctk.set_appearance_mode("light")
    _resolve_fonts()
    boot.title("불러오는 중")
    boot.geometry(f"{int(360 * DPI_SCALE)}x{int(170 * DPI_SCALE)}")
    boot.configure(fg_color=COLORS["bg"])
    boot.attributes("-topmost", True)
    frame = ctk.CTkFrame(boot, fg_color=COLORS["card"], corner_radius=22)
    frame.pack(fill="both", expand=True, padx=16, pady=16)
    ctk.CTkLabel(
        frame,
        text="오늘의 일정을 불러오는 중…",
        font=ui_font(17, emphasize=True),
        text_color=COLORS["dept"],
    ).pack(expand=True)
    boot.update()

    payload = load_today_schedule()
    boot.destroy()

    app = TodaySchedulePopup(payload)
    app.mainloop()


if __name__ == "__main__":
    # 시작 프로그램으로 켜질 때는 네트워크/데스크톱이 준비될 때까지 잠시 기다립니다.
    if "--startup" in sys.argv:
        time.sleep(12)
    show_loading_then_popup()
